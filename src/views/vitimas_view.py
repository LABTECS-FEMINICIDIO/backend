from src.utils.parse_date import parse_to_date
from fastapi import HTTPException, UploadFile
from fastapi.encoders import jsonable_encoder
from sqlalchemy.orm import joinedload
from io import BytesIO
from src.database.database import engine
from src.model.models import VitimasModels
from pydantic import BaseModel
from sqlalchemy.orm import sessionmaker
from uuid import UUID
from typing import Optional
from sqlalchemy import asc, desc
from openpyxl import load_workbook


class Vitima(BaseModel):
    nome: str
    idade: int
    rua: str
    armaUsada: str


class VitimaEdit(BaseModel):
    datadofato: Optional[str] = None
    diah: Optional[str] = None
    horario: Optional[str] = None
    turno: Optional[str] = None
    nome: Optional[str] = None
    idade: Optional[int] = None
    racacor1: Optional[str] = None
    estciv2: Optional[str] = None
    bairro: Optional[str] = None
    rua_beco_travessa_estrada_ramal: Optional[str] = None
    endcomplemento: Optional[str] = None
    tipoarma1: Optional[str] = None
    tipoarma2: Optional[str] = None
    loclesao1: Optional[str] = None
    loclesao2: Optional[str] = None
    loclesao3: Optional[str] = None
    hospitalizacao: Optional[str] = None
    violsexual: Optional[str] = None
    latrocinio: Optional[str] = None
    localdeocorrencia: Optional[str] = None
    presencafilhofamiliar: Optional[str] = None
    compexcomp: Optional[str] = None
    gestacao: Optional[str] = None
    filhosdescrever: Optional[str] = None
    lat: Optional[str] = None
    lng: Optional[str] = None
    sites_in_bulk: Optional[str] = None


async def import_xlsx_file(file: UploadFile):
    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Arquivo inválido. Envie um .xlsx")

    db = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    db_session = db()

    try:
        wb = load_workbook(file.file, data_only=True)
        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            raise HTTPException(status_code=400, detail="Planilha vazia ou sem dados")

        headers = [str(h).strip().lower() if h else None for h in rows[0]]

        valid_columns = {
            col.name.strip().lower()
            for col in VitimasModels.__table__.columns
            if col.name != "id"
        }

        print("VALID COLUMNS", valid_columns)

        header_mapping = {
            "x_lat": "lat",
            "y_long": "lng",
        }

        registros = []

        for row in rows[1:]:
            row_data = {}

            for header, value in zip(headers, row):
                print("header -->", header)
                column_name = header_mapping.get(header, header)
                print("column_name", column_name)
                if not header or column_name not in valid_columns:
                    continue

                if value in ("", None):
                    row_data[column_name] = None
                elif header == "datadofato":
                    row_data[column_name] = parse_to_date(value)
                elif header == "sitegeo1":
                    row_data["siteGeo1"] = value
                elif header == "sitegeo2":
                    row_data["siteGeo2"] = value
                elif header == "sitegeo3":
                    row_data["siteGeo3"] = value
                else:
                    row_data[column_name] = value

            if row_data:
                registros.append(VitimasModels(**row_data))

        if not registros:
            raise HTTPException(
                status_code=400, detail="Nenhum registro válido encontrado"
            )

        db_session.bulk_save_objects(registros)
        db_session.commit()

        return {
            "message": "Importação concluída com sucesso",
            "registros_inseridos": len(registros),
        }

    except Exception as e:
        print(e)
        db_session.rollback()
        raise HTTPException(status_code=500, detail=str(e))

    finally:
        db_session.close()


async def create_vitima(vitima: dict):
    db = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    db_session = db()

    if "lat" in vitima and vitima["lat"] is not None:
        vitima["lat"] = f"{float(vitima['lat']):.5f}"
    if "lng" in vitima and vitima["lng"] is not None:
        vitima["lng"] = f"{float(vitima['lng']):.5f}"

    db_vitima = VitimasModels(**vitima.model_dump())
    db_session.add(db_vitima)
    db_session.commit()
    db_session.refresh(db_vitima)
    return jsonable_encoder(db_vitima)


async def list_vitimas():
    db = sessionmaker(bind=engine)
    db_session = db()
    vitimas = (
        db_session.query(VitimasModels)
        .options(joinedload(VitimasModels.sites))
        .order_by(desc(VitimasModels.datadofato))
        .all()
    )

    return jsonable_encoder(vitimas)


async def list_vitimas_for_export():
    db = sessionmaker(bind=engine)
    db_session = db()
    vitimas = (
        db_session.query(VitimasModels)
        .options(joinedload(VitimasModels.sites))
        .order_by(asc(VitimasModels.idade))
        .all()
    )

    return jsonable_encoder(vitimas)


async def list_one_vitima(vitima_id: UUID):
    db = sessionmaker(bind=engine)
    vitima = db.query(VitimasModels).filter(VitimasModels.id == vitima_id).first()
    if vitima is None:
        raise HTTPException(status_code=404, detail="Vítima not found")
    return jsonable_encoder(vitima)


async def update_vitima(vitima_id: UUID, vitima: dict):
    db = sessionmaker(bind=engine)
    db_session = db()
    db_vitima = (
        db_session.query(VitimasModels).filter(VitimasModels.id == vitima_id).first()
    )

    if db_vitima is None:
        raise HTTPException(status_code=404, detail="Vítima not found")

    for key, value in vitima.items():
        setattr(db_vitima, key, value)

    db_session.commit()
    db_session.refresh(db_vitima)

    return jsonable_encoder(db_vitima)


async def delete_vitima(
    vitima_id: UUID,
):
    db = sessionmaker(bind=engine)
    db_session = db()
    db_vitima = (
        db_session.query(VitimasModels).filter(VitimasModels.id == vitima_id).first()
    )
    if db_vitima is None:
        raise HTTPException(status_code=404, detail="Vítima not found")

    db_session.delete(db_vitima)
    db_session.commit()

    return jsonable_encoder(db_vitima)


async def list_one_vitima(vitima_id: UUID):
    db = sessionmaker(bind=engine)
    db_session = db()
    db_vitima = (
        db_session.query(VitimasModels).filter(VitimasModels.id == vitima_id).first()
    )
    if db_vitima is None:
        raise HTTPException(status_code=404, detail="Vítima not found")

    return jsonable_encoder(db_vitima)


async def delete_all_vitimas():
    db = sessionmaker(bind=engine)()
    try:
        db.query(VitimasModels).delete()
        db.commit()
        return {"mensagem": "Todos os registros foram apagados com sucesso!"}
    except Exception as e:
        db.rollback()
        raise e
