from uuid import UUID
from zoneinfo import ZoneInfo
from fastapi import APIRouter, HTTPException, UploadFile, File
from pydantic import BaseModel
from typing import List, Optional
from src.views.vitimas_view import (
    create_vitima,
    delete_all_vitimas,
    delete_vitima,
    import_xlsx_file,
    list_vitimas,
    update_vitima,
    list_one_vitima,
    list_vitimas_for_export,
)
from src.views.sites_view import list_iml, list_iml_for_export
from fastapi.responses import StreamingResponse
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl import Workbook
from io import BytesIO
from fastapi.responses import FileResponse
from datetime import datetime
import os

router = APIRouter()


class Vitima(BaseModel):
    datadofato: str
    diah: Optional[str] = None
    horario: Optional[str] = None
    turno: Optional[str] = None
    nome: Optional[str] = None
    idade: Optional[int] = None
    racacor1: Optional[str] = None
    estciv2: Optional[str] = None
    bairro: Optional[str] = None
    rua_beco_travessa_estrada_ramal: Optional[str] = None
    endcomplemento: Optional[str] = None
    tipoarma1: Optional[str] = None
    tipoarma2: Optional[str] = None
    loclesao1: Optional[str] = None
    loclesao2: Optional[str] = None
    loclesao3: Optional[str] = None
    hospitalizacao: Optional[str] = None
    compexcomp: Optional[str] = None
    lat: Optional[str] = None
    lng: Optional[str] = None
    violsexual: Optional[str] = None
    latrocinio: Optional[str] = None
    site1: Optional[str] = None
    site2: Optional[str] = None
    site3: Optional[str] = None
    siteGeo1: Optional[str] = None
    siteGeo2: Optional[str] = None
    siteGeo3: Optional[str] = None
    zona: Optional[str] = None
    localdeocorrencia: Optional[str] = None
    presencafilhofamiliar: Optional[str] = None
    sites_in_bulk: Optional[str] = None
    gestacao: Optional[str] = None
    filhosdescrever: Optional[int] = None
    sites: Optional[List[UUID]] = []


@router.post("/vitimas/")
async def create_vitimas_controller(vitima: Vitima):
    return await create_vitima(vitima)


@router.post("/import-xlsx")
async def import_xlsx(file: UploadFile = File(...)):
    return await import_xlsx_file(file)


@router.get("/vitimas/{vitima_id}")
async def list_one_controller(vitima_id: str):
    return await list_one_vitima(vitima_id)


@router.get("/vitimas/")
async def list_tags_controller():
    return await list_vitimas()


@router.get("/delete/vitimas")
async def delete_vitimas():
    return await delete_all_vitimas()


# @router.get("/tag/{tag_name}", response_model=Tag)
# async def list_one_tag_controller(tag_name: str):
#     return await list_one_tag(tag_name)


@router.patch("/vitimas/{vitima_id}")
async def update_vitima_controller(vitima_id, item: dict):
    return await update_vitima(vitima_id, item)


@router.delete("/vitimas/{vitima_id}")
async def delete_vitimas_controller(vitima_id: str):
    return await delete_vitima(vitima_id)


headers = [
    "numero1",
    "registro_oficial_do",
    "naocapturado",
    "homicidio",
    "numerodo",
    "datadofato",
    "diah",
    "diasemh",
    "mesh",
    "anoh",
    "horario",
    "turno",
    "nome",
    "idade",
    "racacor1",
    "estciv2",
    "esc2",
    "bairro",
    "zona",
    "rua_beco_travessa_estrada_ramal",
    "endcomplemento",
    "local_desova_corpo",
    "X_lat",
    "Y_long",
    "precisao_local_classificacao",
    "coordenadas_derivam",
    "cid10cod4final",
    "cid10cod4finaltexto",
    "tipoarma1",
    "tipoarma2",
    "loclesao1",
    "loclesao2",
    "loclesao3",
    "localdaslesoes",
    "numerodelesoes",
    "possivelfemin",
    "discord_classificacoes",
    "possivelfemin1",
    "hospitalizacao",
    "vitusudrogilicita",
    "relacaotraf",
    "violsexual",
    "usodealcool",
    "latrocinio",
    "tipoviol",
    "localdeocorrencia",
    "presencafilhofamiliar",
    "situacaorua",
    "nivsupcom",
    "represalia trafico",
    "sexoagressor",
    "compexcomp",
    "outrofamconhefam",
    "compexfam",
    "excompan",
    "conhecido",
    "circunsmorte",
    "gestacao",
    "puerperio",
    "filhosdescrever",
    "menor14anos",
    "maior60anos",
    "vulnerabil_fisica_mental",
    "presenca_ascend_descendente",
    "presenca_medida_protet_urgen",
    "tipo_intimo_naointimo",
    "inf_bol_ocorrencia_iml_bol",
    "inf_bol_ocorrencia_revisao",
    "consulta_jud_bol1",
    "consulta_jud_bol2",
    "consulta_jud_revisao1",
    "consulta_jud_revisao2",
    "observacoes",
    "SITE1",
    "SITE2",
    "SITE3",
    "SITEGEO1",
    "SITEGEO2",
    "SITEGEO3",
    "check_30dias",
    "data_versao_do",
    "cidade",
]

dictionary = {
    "numero1": "numero1",
    "registro_oficial_do": "registro_fvs_do",
    "naocapturado": "naocapturado",
    "homicidio": "homicidio",
    "numerodo": "numerodo",
    "datadofato": "datadofato",
    "diah": "diah",
    "diasemh": "diasemh",
    "mesh": "mesh",
    "anoh": "anoh",
    "horario": "horario",
    "turno": "turno",
    "nome": "nome",
    "idade": "idade",
    "racacor1": "racacor1",
    "estciv2": "estciv2",
    "esc2": "esc2",
    "bairro": "bairro",
    "zona": "zona",
    "rua_beco_travessa_estrada_ramal": "rua_beco_travessa_estrada_ramal",
    "endcomplemento": "endcomplemento",
    "local_desova_corpo": "local_desova_corpo",
    "X_lat": "X_Lat",
    "Y_long": "Y_Long",
    "precisao_local_classificacao": "precisao_local_classificacao",
    "coordenadas_derivam": "coordenadas_derivam",
    "cid10cod4final": "cid10cod4final",
    "cid10cod4finaltexto": "cid10cod4finaltexto",
    "tipoarma1": "tipoarma1",
    "tipoarma2": "tipoarma2",
    "loclesao1": "loclesao1",
    "loclesao2": "loclesao2",
    "loclesao3": "loclesao3",
    "localdaslesoes": "localdaslesoes",
    "numerodelesoes": "numerodelesoes",
    "possivelfemin": "possivelfemin",
    "discord_classificacoes": "discord_classificacoes",
    "possivelfemin1": "possivelfemin1",
    "hospitalizacao": "hospitalizacao",
    "vitusudrogilicita": "vitusudrogilicita",
    "relacaotraf": "relacaotraf",
    "violsexual": "violsexual",
    "usodealcool": "usodealcool",
    "latrocinio": "latrocinio",
    "tipoviol": "tipoviol",
    "localdeocorrencia": "localdeocorrencia",
    "presencafilhofamiliar": "presencafilhofamiliar",
    "situacaorua": "situacaorua",
    "nivsupcom": "nivsupcomouincomp",
    "represalia trafico": "represalia trafico",
    "sexoagressor": "sexoagressor",
    "compexcomp": "compexcomp",
    "outrofamconhefam": "outrofamconhefam",
    "compexfam": "compexfam",
    "excompan": "excompan",
    "conhecido": "conhecido",
    "circunsmorte": "circunsmorte",
    "gestacao": "gestacao",
    "puerperio": "puerperio",
    "filhosdescrever": "filhosdescrever",
    "menor14anos": "menor14anos",
    "maior60anos": "maior60anos",
    "vulnerabil_fisica_mental": "vulnerabil_fisica_mental",
    "presenca_ascend_descendente": "presenca_ascend_descendente",
    "presenca_medida_protet_urgen": "presenca_medida_protet_urgen",
    "tipo_intimo_naointimo": "tipo_intimo_naointimo",
    "inf_bol_ocorrencia_iml_bol": "inf_bol_ocorrencia_iml_bol",
    "inf_bol_ocorrencia_revisao": "inf_bol_ocorrencia_revisao",
    "consulta_jud_bol1": "consulta_saj_bol1",
    "consulta_jud_bol2": "consulta_saj_bol2",
    "consulta_jud_revisao1": "consulta_saj_revisao1",
    "consulta_jud_revisao2": "consulta_saj_revisao2",
    "observacoes": "observacoes",
    "sites_in_bulk": "sites_in_bulk",
    "SITE1": "site1",
    "SITE2": "site2",
    "SITE3": "site3",
    "SITEGEO1": "siteGeo1",
    "SITEGEO2": "siteGeo2",
    "SITEGEO3": "siteGeo3",
    "check_30dias": "check_30dias",
    "data_versao_do": "data_versao_do",
    "cidade": "cidade",
}


def export_to_xlsx(data):
    wb = Workbook()
    ws = wb.active

    ws.append(headers)

    for row in data:

        row_dict = dict(row)
        row_dict["X_Lat"] = row_dict.pop("lat", "NA")
        row_dict["Y_Long"] = row_dict.pop("lng", "NA")

        new_row_dict = {}
        for key, value in row_dict.items():

            if value is None or value == "":
                new_row_dict[key] = "NA"
            elif key == "datadofato":
                try:
                    data_datetime = datetime.fromisoformat(value.replace("Z", "+00:00"))
                    # Formatar a data no formato desejado
                    data_formatada = data_datetime.strftime("%d/%m/%Y")
                    new_row_dict[key] = data_formatada
                except ValueError:
                    new_row_dict[key] = "Invalid date"
            elif key == "zona":
                new_row_dict[key] = value.lower().replace(" ", "")

            else:
                new_row_dict[key] = value

        if "sites" in row_dict:
            for index, site in enumerate(row_dict["sites"]):
                if index < 3:
                    new_row_dict[f"SITE{index + 1}"] = site["link"]
                else:
                    break

        row_data = [new_row_dict.get(dictionary[header], "NA") for header in headers]

        ws.append(row_data)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return output


headers_iml = [
    "dataEntrada",
    "horaEntrada",
    "sexo",
    "idade",
    "bairroDaRemocao",
    "causaMorte",
    "DataCaptura",
    "HoraCaptura",
    "cidade",
]

trad_iml = {
    "dataEntrada": "dataEntrada",
    "horaEntrada": "horaEntrada",
    "sexo": "sexo",
    "idade": "idade",
    "bairroDaRemocao": "bairroDaRemocao",
    "causaMorte": "causaMorte",
    "DataCaptura": "createdAt",
    "HoraCaptura": "createdAt",
    "cidade": "cidade",
}


def export_to_xlsx_iml(data):
    wb = Workbook()
    ws = wb.active

    ws.append(headers_iml)

    for col_index, header_name in enumerate(headers_iml, start=1):
        cell = ws.cell(row=1, column=col_index)

        novo_texto = {
            "dataEntrada": "Data da Entrada",
            "horaEntrada": "Hora da Entrada",
            "sexo": "Sexo",
            "idade": "Idade",
            "bairroDaRemocao": "Bairro da Remoção",
            "causaMorte": "Causa da Morte",
            "DataCaptura": "Data da Captura",
            "HoraCaptura": "Hora da Captura",
            "cidade": "Cidade",
        }.get(header_name, header_name)

        cell.value = novo_texto

        cell.fill = PatternFill(
            start_color="EDE1FE", end_color="EDE1FE", fill_type="solid"
        )
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = Font(bold=True, color="000000")

    for row in data:
        row_dict = dict(row)
        row_data = []

        for header in headers_iml:

            if header == "dataEntrada":
                created_at = datetime.strptime(
                    row_dict["createdAt"].split("+")[0], "%Y-%m-%d %H:%M:%S.%f"
                )
                row_data.append(created_at.date())

            elif header == "horaEntrada":
                created_at = datetime.strptime(
                    row_dict["createdAt"].split("+")[0], "%Y-%m-%d %H:%M:%S.%f"
                )
                created_at = created_at.replace(tzinfo=ZoneInfo("UTC"))
                cliente_tz = ZoneInfo(os.getenv("TZ"))
                created_at_local = created_at.astimezone(cliente_tz)
                row_data.append(created_at_local.strftime("%H:%M:%S"))

            elif header == "DataCaptura":
                created_at = datetime.strptime(
                    row_dict["createdAt"].split("+")[0], "%Y-%m-%d %H:%M:%S.%f"
                )
                row_data.append(created_at.date())

            elif header == "HoraCaptura":
                created_at = datetime.strptime(
                    row_dict["createdAt"].split("+")[0], "%Y-%m-%d %H:%M:%S.%f"
                )
                created_at = created_at.replace(tzinfo=ZoneInfo("UTC"))
                cliente_tz = ZoneInfo(os.getenv("TZ"))
                created_at_local = created_at.astimezone(cliente_tz)
                row_data.append(created_at_local.strftime("%H:%M:%S"))

            elif header == "cidade":
                row_data.append(os.getenv("CITY"))

            else:
                row_data.append(row_dict.get(trad_iml[header], ""))

        ws.append(row_data)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


@router.get("/export-xlsx")
async def export_xlsx():
    try:
        vitimas_data = await list_vitimas_for_export()

        output = export_to_xlsx(vitimas_data)

        return StreamingResponse(
            content=output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=vitimas.xlsx"},
        )

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/iml/export-xlsx")
async def export_iml_xlsx():
    try:
        iml_data = await list_iml_for_export()

        output = export_to_xlsx_iml(iml_data)

        return StreamingResponse(
            content=output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": "attachment; filename=iml.xlsx"},
        )

    except Exception as e:
        print(e)
        raise HTTPException(status_code=500, detail=str(e))
